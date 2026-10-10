"""Validate controlled migration cases on a selected public UCI retail extract.

CSV needs no dependencies. XLSX is read-only with optional openpyxl. Source line
values are NOT bank cash movements; candidate faults are deliberately injected.
"""
import argparse
from copy import deepcopy
import csv
from datetime import datetime, timedelta, timezone
from decimal import Decimal, DecimalException, Inexact, Rounded, Underflow, Overflow, localcontext
import hashlib
from itertools import islice
import json
import re
from pathlib import Path
import sys
from urllib.request import urlopen
import zipfile
from zoneinfo import ZoneInfo
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from cutoverguard.engine import reconcile, canonical

SOURCE_URL = "https://archive.ics.uci.edu/dataset/352/online+retail"
DOWNLOAD_URL = "https://archive.ics.uci.edu/static/public/352/online%2Bretail.zip"
REQUIRED = {"InvoiceNo", "Quantity", "InvoiceDate", "UnitPrice", "CustomerID"}


def text(value):
    if isinstance(value, datetime):
        return value.isoformat()
    return "" if value is None else str(value)


def london_time(value, date_format="%m/%d/%Y %H:%M"):
    if isinstance(value, datetime):
        parsed = value
    else:
        value = text(value)
        if re.search(r"[.,][0-9]{7,}", value):
            raise ValueError("InvoiceDate precision exceeds six fractional digits")
        try:
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            try:
                parsed = datetime.strptime(value, date_format)
            except ValueError as exc:
                raise ValueError("InvoiceDate is invalid; use ISO or the declared CSV --date-format") from exc
    if parsed.tzinfo is not None and parsed.utcoffset() is not None:
        return parsed.astimezone(timezone.utc)
    london = ZoneInfo("Europe/London")
    choices = []
    for fold in (0, 1):
        aware = parsed.replace(tzinfo=london, fold=fold)
        utc = aware.astimezone(timezone.utc)
        if utc.astimezone(london).replace(tzinfo=None) == parsed and utc not in choices:
            choices.append(utc)
    if len(choices) != 1:
        raise ValueError("InvoiceDate is ambiguous or nonexistent in assumed Europe/London timezone")
    return choices[0]


def invoice_minor(quantity, price):
    try:
        if len(text(quantity)) > 64 or len(text(price)) > 64:
            raise ValueError("Quantity/UnitPrice exceeds numeric input length")
        q, p = Decimal(text(quantity)), Decimal(text(price))
        if not q.is_finite() or not p.is_finite():
            raise ValueError("Quantity/UnitPrice must be finite")
        if q != q.to_integral_value() or abs(q) > Decimal("1e12"):
            raise ValueError("Quantity must be an integer with absolute value at most 1e12")
        if p < 0 or p > Decimal("1e12"):
            raise ValueError("UnitPrice must be nonnegative and at most 1e12")
        with localcontext() as ctx:
            ctx.prec = 100
            for signal in (Inexact, Rounded, Underflow, Overflow):
                ctx.traps[signal] = True
            minor = q * p * 100
        if minor != minor.to_integral_value():
            raise ValueError("Quantity × UnitPrice has fractional pennies; no rounding was applied")
        return str(int(minor))
    except (DecimalException, OverflowError) as exc:
        raise ValueError("Quantity/UnitPrice is not a supported exact decimal") from exc


def csv_rows(path):
    with Path(path).open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        if not REQUIRED.issubset(reader.fieldnames or []):
            raise ValueError("CSV requires InvoiceNo, Quantity, InvoiceDate, UnitPrice, CustomerID headers")
        for record_number, row in enumerate(reader, 2):
            yield record_number, row


def xlsx_rows(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("XLSX reading needs openpyxl. Run .venv/bin/python -m pip install openpyxl, or supply an exported CSV with --csv.") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = [text(value) for value in next(rows)]
        if not REQUIRED.issubset(headers):
            raise ValueError("Workbook requires the original UCI Online Retail headers")
        for record_number, row in enumerate(rows, 2):
            yield record_number, dict(zip(headers, row))
    finally:
        workbook.close()


def build_bundle(rows, *, source_sha256, captured_at, limit=2000, date_format="%m/%d/%Y %H:%M", input_name="public extract"):
    if type(limit) is not int or not 1 <= limit <= 4999:
        raise ValueError("--limit must be 1–4999 scanned records (one slot reserved for a duplicate test)")
    events, observations, excluded, accounts, times = [], [], [], set(), []
    scanned = 0
    for row_number, raw in islice(rows, limit):
        scanned += 1
        preserved = {str(key): text(value) for key, value in raw.items()}
        try:
            customer = text(raw.get("CustomerID"))
            if not customer.strip():
                raise ValueError("CustomerID is missing; no customer/account was invented")
            invoice = text(raw.get("InvoiceNo"))
            if not invoice.strip():
                raise ValueError("InvoiceNo is missing")
            amount = invoice_minor(raw.get("Quantity"), raw.get("UnitPrice"))
            timestamp = london_time(raw.get("InvoiceDate"), date_format)
            account = "receivable:" + customer
            event = "uci-row-" + str(row_number).zfill(9)
            events.append(dict(event_id=event, account_id=account, amount_minor=amount, currency="GBP", event_time=timestamp.isoformat()))
            observations.append(dict(original_record_number=row_number, event_id=event, raw_record=preserved))
            accounts.add(customer)
            times.append(timestamp)
        except (ValueError, TypeError, OverflowError) as exc:
            excluded.append(dict(original_record_number=row_number, reason=str(exc), raw_record=preserved))
    if not events:
        raise ValueError("No validated invoice lines in selected extract; inspect dates, customers and decimal amounts")
    try:
        start, end = min(times), max(times) + timedelta(microseconds=1)
    except OverflowError as exc:
        raise ValueError("Selected event timestamps leave no supported UTC half-open window end") from exc
    snapshot = dict(snapshot_id="uci-selected-" + source_sha256[:16], captured_at=captured_at,
                    mapping_version="customer-receivable-adapter-v1", window_start=start.isoformat(), window_end=end.isoformat(),
                    complete_through=end.isoformat(), profiles=[dict(account_id="receivable:" + customer, customer_id=customer) for customer in sorted(accounts)], transactions=events)
    snapshot["snapshot_id"] = "uci-selected-" + hashlib.sha256(canonical({"source_file_sha256": source_sha256, "adapter_version": "1.0", "date_format": date_format, "selected_snapshot": snapshot}).encode()).hexdigest()
    dataset = dict(name="UCI Online Retail", source_url=SOURCE_URL, license="CC BY 4.0", source_file_sha256=source_sha256,
                   validation_kind="historical_source_controlled_candidate",
                   description=f"Historical public retail invoice-line values with deliberately controlled candidate defects. Selected {len(events)} validated lines from {scanned} scanned records; {len(excluded)} excluded. Customer receivables buckets are not bank accounts; this is not observed cash, settlement or a real migration.",
                   scanned_records=scanned, selected_records=len(events), excluded_records=len(excluded),
                   extraction_counts={"scanned_records": scanned, "selected_records": len(events), "excluded_records": len(excluded)},
                   financial_measure="quantity_times_unit_price_invoice_line_value", account_semantics="customer-derived receivables buckets; not bank accounts",
                   completeness_scope="selected validated extract only; not complete retailer or bank ledger",
                   timezone_assumption="Naive InvoiceDate assumed Europe/London; ambiguous/nonexistent DST excluded. Aware dates retain their declared offset.",
                   capture_semantics="captured_at is supplied file-modification timestamp, not an independently verified system snapshot capture",
                   csv_date_format=date_format, input_name=input_name)
    bundle = dict(schema_version="1.0", data_kind="historical", config={"rule_version": "1.0", "currency_scales": {"GBP": 2}},
                  source=snapshot, candidate=deepcopy(snapshot), dataset_info=dataset)
    bundle["candidate"]["snapshot_id"] += "-identity-control"
    manifest = dict(data_kind="historical", validation_kind="historical_source_controlled_candidate", dataset_info=deepcopy(dataset),
                    scanned_record_limit=limit, scanned_records=scanned, selected_records=len(events), excluded_records=len(excluded),
                    unseen_records="not counted; records beyond scan limit are outside this extract",
                    exclusions=excluded, selected_observations=observations,
                    selection_policy="First N original logical records in file order; no random sampling, no inferred customers, no rounding.",
                    completeness="Only the selected valid records are compared; excluded/unscanned source records remain outside reconciliation coverage.")
    return bundle, manifest


def controlled_cases(baseline):
    source = baseline["source"]["transactions"]
    if len(source) < 2:
        raise ValueError("Controlled validation needs at least two valid invoice lines")
    first, second = source[0], source[1]
    distinct = next((row for row in source[1:] if row["account_id"] != first["account_id"]), None)
    if distinct is None:
        raise ValueError("Controlled ownership test needs at least two distinct customers; increase --limit")
    cases = []
    def add(name, bundle, expected):
        bundle["candidate"]["snapshot_id"] = baseline["source"]["snapshot_id"] + "-controlled-" + name
        bundle["dataset_info"]["controlled_case"] = name
        cases.append((name, bundle, set(expected)))
    add("baseline", deepcopy(baseline), [])
    b = deepcopy(baseline)
    left = next(r for r in b["candidate"]["transactions"] if r["event_id"] == first["event_id"])
    right = next(r for r in b["candidate"]["transactions"] if r["event_id"] == distinct["event_id"])
    left["account_id"], right["account_id"] = right["account_id"], left["account_id"]
    add("ownership_swaps", b, [("account_mismatch", first["event_id"]), ("account_mismatch", distinct["event_id"])])
    b = deepcopy(baseline)
    b["candidate"]["transactions"].pop(0)
    extra = deepcopy(first)
    extra["event_id"] = "controlled-extra-event"
    b["candidate"]["transactions"].append(extra)
    add("missing_extra", b, [("event_missing", first["event_id"]), ("event_extra", extra["event_id"])])
    b = deepcopy(baseline)
    b["candidate"]["transactions"][0]["amount_minor"] = str(int(first["amount_minor"]) + 1)
    b["candidate"]["transactions"][1]["amount_minor"] = str(int(second["amount_minor"]) - 1)
    add("compensating_penny_changes", b, [("amount_mismatch", first["event_id"]), ("amount_mismatch", second["event_id"])])
    b = deepcopy(baseline)
    b["candidate"]["transactions"].append(deepcopy(first))
    b["candidate"]["profiles"].append(deepcopy(b["candidate"]["profiles"][0]))
    add("duplicate_keys", b, [("duplicate_event_key", first["event_id"]), ("duplicate_profile_key", b["candidate"]["profiles"][0]["account_id"])])
    b = deepcopy(baseline)
    profile = b["candidate"]["profiles"][0]
    profile["customer_id"] = "controlled-mapping-customer"
    add("customer_mapping", b, [("customer_mismatch", profile["account_id"])])
    return cases


def validate_cases(baseline, output=None):
    results = []
    for name, bundle, expected in controlled_cases(baseline):
        report = reconcile(bundle)
        actual_list = [(i["type"], i["key"]) for i in report["issues"] if i["severity"] == "critical" and i["check_id"] != "business"]
        actual = set(actual_list)
        success = actual == expected and len(actual_list) == len(expected) and report["status"] == ("ready_for_review" if name == "baseline" else "blocked")
        results.append(dict(case=name, expected_primary=[list(item) for item in sorted(expected)], actual_primary=[list(item) for item in sorted(actual)],
                            expected_primary_count=len(expected), actual_primary_count=len(actual_list), verified=success, review_status=report["status"], run_id=report["run_id"]))
        if output:
            target = Path(output)
            target.mkdir(parents=True, exist_ok=True)
            (target / (name + "-bundle.json")).write_text(json.dumps(bundle, indent=2) + "\n")
            (target / (name + "-report.json")).write_text(json.dumps(report, indent=2) + "\n")
    return dict(validation_kind="historical_source_controlled_candidate", data_kind="historical", source_file_sha256=baseline["dataset_info"]["source_file_sha256"],
                status="verified" if all(row["verified"] for row in results) else "failed", cases=results,
                limitations=["These are deliberately injected candidate changes, not observed migration failures.", "Invoice-line values are not bank balances, cash movements or settlement amounts.", "Only the selected validated extract is covered; excluded/unscanned rows are disclosed separately.", "Local validation does not establish Databricks execution or real-world detection accuracy."])


def download_workbook(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    archive = folder / "online-retail.zip"
    try:
        with urlopen(DOWNLOAD_URL, timeout=30) as response, archive.open("wb") as handle:
            total = 0
            while chunk := response.read(1024 * 1024):
                total += len(chunk)
                if total > 40_000_000:
                    raise ValueError("Official download exceeds 40 MB safety limit")
                handle.write(chunk)
        with zipfile.ZipFile(archive) as zipped:
            name = next((n for n in zipped.namelist() if n.lower().endswith(".xlsx")), None)
            if name is None:
                raise ValueError("Official archive contains no XLSX workbook")
            if zipped.getinfo(name).file_size > 80_000_000:
                raise ValueError("Workbook exceeds 80 MB extraction limit")
            path = folder / "Online Retail.xlsx"
            with zipped.open(name) as source, path.open("wb") as dest:
                while chunk := source.read(1024 * 1024):
                    dest.write(chunk)
            return path
    except (OSError, zipfile.BadZipFile) as exc:
        raise ValueError("Official download failed. Download the UCI archive in your browser and use --xlsx FILE, or export its sheet and use --csv FILE.") from exc


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--csv", type=Path)
    group.add_argument("--xlsx", type=Path)
    group.add_argument("--download", action="store_true")
    parser.add_argument("--output", type=Path, default=Path("public_retail_validation"))
    parser.add_argument("--limit", type=int, default=2000)
    parser.add_argument("--date-format", default="%m/%d/%Y %H:%M")
    args = parser.parse_args()
    try:
        if args.download:
            import importlib.util
            if importlib.util.find_spec("openpyxl") is None:
                raise ValueError("--download needs optional openpyxl to read the official XLSX. Run .venv/bin/python -m pip install openpyxl, then repeat; or use --csv with no extra dependency.")
        path = download_workbook(args.output / "download") if args.download else args.csv or args.xlsx
        if not path.is_file():
            raise ValueError("Input file does not exist")
        hasher = hashlib.sha256()
        with path.open("rb") as handle:
            while chunk := handle.read(1024 * 1024):
                hasher.update(chunk)
        rows = csv_rows(path) if args.csv else xlsx_rows(path)
        captured = datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat()
        bundle, manifest = build_bundle(rows, source_sha256=hasher.hexdigest(), captured_at=captured, limit=args.limit, date_format=args.date_format, input_name=path.name)
        summary = validate_cases(bundle, args.output)
        args.output.mkdir(parents=True, exist_ok=True)
        (args.output / "extract-manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
        (args.output / "validation-summary.json").write_text(json.dumps(summary, indent=2) + "\n")
        print(json.dumps({"status": summary["status"], "scanned_records": manifest["scanned_records"], "selected_records": manifest["selected_records"], "excluded_records": manifest["excluded_records"], "controlled_cases": len(summary["cases"]), "output": str(args.output), "scope": "Historical invoice-line extract; deliberately injected candidate defects; local execution"}, indent=2))
        return 0 if summary["status"] == "verified" else 1
    except (ValueError, OSError) as exc:
        print("Public retail validation: " + str(exc), file=sys.stderr)
        return 2

if __name__ == "__main__":
    raise SystemExit(main())
